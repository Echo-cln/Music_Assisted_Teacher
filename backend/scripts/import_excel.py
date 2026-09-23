import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import delete, func, select

BACKEND_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_DIR))

from app.db.init_db import init_db  # noqa: E402
from app.db.session import SessionLocal  # noqa: E402
from app.models.entities import MusicTheory, Song, TeachingGame, TeachingMistake  # noqa: E402

REGION_SHEETS = ["华南地区", "西南地区", "西北地区", "华中地区", "华东地区", "华北地区", "东北地区"]


def value(row, index, default=""):
    current = row[index] if index < len(row) else None
    return default if current is None else current


def import_workbook(path: Path, replace: bool = False) -> dict[str, int]:
    init_db()
    workbook = load_workbook(path, read_only=True, data_only=True)
    with SessionLocal() as db:
        if replace:
            for model in (Song, TeachingGame, MusicTheory, TeachingMistake):
                db.execute(delete(model))
            db.commit()
        if db.scalar(select(func.count()).select_from(Song)):
            return {"songs": 0, "games": 0, "theory": 0, "mistakes": 0}

        counts = {"songs": 0, "games": 0, "theory": 0, "mistakes": 0}
        for sheet_name in REGION_SHEETS:
            sheet = workbook[sheet_name]
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if not value(row, 1):
                    continue
                db.add(
                    Song(
                        source_row=int(value(row, 0, 0)),
                        name=str(value(row, 1)),
                        province=str(value(row, 2)),
                        mood=str(value(row, 3)),
                        mode=str(value(row, 4)),
                        grade=str(value(row, 5)),
                        source=str(value(row, 6)),
                        song_type=str(value(row, 7)),
                        range_note=str(value(row, 8)),
                        range_score=int(value(row, 9, 1)),
                        rhythm_score=int(value(row, 10, 1)),
                        dialect_score=int(value(row, 11, 1)),
                        difficulty=str(value(row, 12, "1星")),
                        region=sheet_name,
                    )
                )
                counts["songs"] += 1

        for row in workbook["课前游戏知识库"].iter_rows(min_row=2, values_only=True):
            if value(row, 1):
                db.add(
                    TeachingGame(
                        category=str(value(row, 0)),
                        name=str(value(row, 1)),
                        personality=str(value(row, 2)),
                        grade=str(value(row, 3)),
                        match_condition=str(value(row, 4)),
                        instructions=str(value(row, 5)),
                    )
                )
                counts["games"] += 1

        for row in workbook["乐理知识讲解"].iter_rows(min_row=2, values_only=True):
            if value(row, 1):
                db.add(
                    MusicTheory(
                        category=str(value(row, 0)),
                        term=str(value(row, 1)),
                        lower_grade_script=str(value(row, 2)),
                        upper_grade_script=str(value(row, 3)),
                    )
                )
                counts["theory"] += 1

        for row in workbook["易错点预判与纠正库"].iter_rows(min_row=2, values_only=True):
            if value(row, 0):
                db.add(
                    TeachingMistake(
                        category=str(value(row, 0)), problem=str(value(row, 1)), correction=str(value(row, 2))
                    )
                )
                counts["mistakes"] += 1

        db.commit()
        return counts


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="将音乐教学 Excel 数据导入 SQLite")
    parser.add_argument("--file", type=Path, default=BACKEND_DIR / "data" / "raw" / "music_resources.xlsx")
    parser.add_argument("--replace", action="store_true")
    args = parser.parse_args()
    result = import_workbook(args.file, replace=args.replace)
    print("导入完成：", result)
